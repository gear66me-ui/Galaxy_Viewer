from pathlib import Path
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src = Path("audit/distance-age-review/galaxy-distance-age-source-audit.csv")
dst = Path("audit/distance-age-review/galaxy-distance-age-source-audit.xlsx")
if not src.exists():
    raise SystemExit(f"CSV not found: {src}")

wb = Workbook(write_only=False)
ws = wb.active
ws.title = "Distance Audit"
with src.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    headers = next(reader)
    ws.append(headers)
    count = 0
    widths = [min(max(12, len(h) + 2), 42) for h in headers]
    for row in reader:
        ws.append(row)
        count += 1
        for i, val in enumerate(row[:len(widths)]):
            if val:
                widths[i] = min(max(widths[i], min(len(str(val)) + 2, 42)), 42)

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 32
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
for i, width in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"
wb.properties.title = "Galaxy Viewer Distance Source Audit"
wb.properties.subject = "SIMBAD and NED distance review template"
wb.properties.creator = "Galaxy Viewer"
wb.save(dst)
print(f"Created {dst} with {count:,} records and {len(headers)} columns.")
